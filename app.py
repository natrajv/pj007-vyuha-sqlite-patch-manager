import os
import sqlite3
import uuid
import shutil
import json
import math
import fnmatch
from datetime import date, datetime, time
from pathlib import Path
from flask import Flask, render_template, request, redirect, url_for, send_file, session, flash
from werkzeug.utils import secure_filename
from openpyxl import load_workbook

from merger.comparator import (
    compare_databases, compare_schema, get_table_changes, get_table_preview,
    quote_ident, tables,
)
from merger.merger import merge_databases

BASE_DIR = Path(__file__).resolve().parent
WORKSPACE = BASE_DIR / "workspace"
WORKSPACE.mkdir(exist_ok=True)

app = Flask(__name__)
app.secret_key = "change-this-secret-key"
app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024

CONFIG_PATH = BASE_DIR / "config.json"
with CONFIG_PATH.open("r", encoding="utf-8") as f:
    APP_CONFIG = json.load(f)

APP_NAME = "Vyuha SQLite Patch Manager"
APP_VERSION = str(APP_CONFIG.get("version", "1.0.0"))
PAGINATION_SIZE = max(1, int(APP_CONFIG.get("pagination_size", 10)))


@app.context_processor
def inject_app_config():
    return {
        "app_name": APP_NAME,
        "app_version": APP_VERSION,
    }


def db_paths():
    sid = session.get("session_id")
    if not sid:
        return None, None, None
    d = WORKSPACE / sid
    return d, d / session.get("patch_db_file", "source.sqlite"), d / "target.sqlite"


def create_patch_from_excel(workbook_path, base_path, output_path):
    """Build a patch database from workbook sheets matching Base tables."""
    with sqlite3.connect(base_path) as base:
        base_tables = {
            name.casefold(): name
            for (name,) in base.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"
            )
        }
        table_columns = {}
        for name in base_tables.values():
            table_columns[name] = base.execute(
                f'PRAGMA table_info("{name.replace(chr(34), chr(34) * 2)}")'
            ).fetchall()

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    matches = []
    matched_tables = set()
    try:
        for sheet in workbook.worksheets:
            table_name = base_tables.get(sheet.title.casefold())
            if not table_name or table_name in matched_tables:
                continue
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                continue
            header_map = {}
            for index, header in enumerate(headers):
                if header is not None and str(header).strip():
                    header_map.setdefault(str(header).strip().casefold(), index)
            base_cols = table_columns[table_name]
            selected = [
                (column, header_map[column[1].casefold()])
                for column in base_cols if column[1].casefold() in header_map
            ]
            if not selected:
                continue
            sheet_rows = []
            for row in rows:
                values = [row[index] if index < len(row) else None for _, index in selected]
                if any(value is not None for value in values):
                    sheet_rows.append([
                        value.isoformat() if isinstance(value, (date, datetime, time)) else
                        int(value) if isinstance(value, bool) else value
                        for value in values
                    ])
            matches.append((table_name, selected, sheet_rows))
            matched_tables.add(table_name)
    finally:
        workbook.close()

    if not matches:
        raise ValueError("The workbook must contain a sheet named for a Base table and at least one matching column.")

    with sqlite3.connect(output_path) as patch:
        for table_name, selected, rows in matches:
            definitions = [
                f'{quote_ident(column[1])} {column[2] or ""}'.rstrip()
                for column, _ in selected
            ]
            selected_pk = [(column[5], column[1]) for column, _ in selected if column[5]]
            all_pk = [column[1] for column in table_columns[table_name] if column[5]]
            if selected_pk and len(selected_pk) == len(all_pk):
                pk_names = [name for _, name in sorted(selected_pk)]
                definitions.append(f'PRIMARY KEY ({", ".join(quote_ident(name) for name in pk_names)})')
            patch.execute(f'CREATE TABLE {quote_ident(table_name)} ({", ".join(definitions)})')
            if rows:
                placeholders = ",".join("?" for _ in selected)
                patch.executemany(
                    f'INSERT INTO {quote_ident(table_name)} VALUES ({placeholders})', rows
                )


@app.route("/")
def index():
    stats = session.get("stats")
    return render_template("index.html", stats=stats)


@app.post("/upload")
def upload():
    source = request.files.get("source_db")
    target = request.files.get("target_db")

    if not source or not target or not source.filename or not target.filename:
        flash("Please select both a Patch SQLite/Excel file and a Base SQLite database.", "danger")
        return redirect(url_for("index"))

    is_excel = source.filename.lower().endswith((".xlsx", ".xlsm"))
    if not is_excel and not source.filename.lower().endswith((".db", ".sqlite", ".sqlite3")):
        flash("Patch file must be a SQLite database or Excel workbook (.xlsx/.xlsm).", "danger")
        return redirect(url_for("index"))

    if not target.filename.lower().endswith((".db", ".sqlite", ".sqlite3")):
        flash("Base SQLite database must be a SQLite database file.", "danger")
        return redirect(url_for("index"))

    sid = uuid.uuid4().hex
    d = WORKSPACE / sid
    d.mkdir(parents=True, exist_ok=True)

    source_path = d / "source.sqlite"
    target_path = d / "target.sqlite"
    source_upload_path = d / secure_filename(source.filename)
    target.save(target_path)

    try:
        source.save(source_upload_path)
        if is_excel:
            generated_name = f"{Path(secure_filename(source.filename)).stem}.sqlite"
            source_path = d / generated_name
            create_patch_from_excel(source_upload_path, target_path, source_path)
        else:
            source_upload_path.replace(source_path)
        stats = compare_databases(source_path, target_path)
        schema_result = compare_schema(source_path, target_path)
    except Exception as exc:
        shutil.rmtree(d, ignore_errors=True)
        flash(f"Database comparison failed: {exc}", "danger")
        return redirect(url_for("index"))

    session["session_id"] = sid
    session["stats"] = stats
    session["schema_match"] = schema_result.get("matches", False)
    session["schema_report"] = None
    session["patch_name"] = source.filename
    session["patch_db_file"] = source_path.name
    session["base_name"] = target.filename
    session.pop("merged", None)
    return redirect(url_for("dashboard"))


@app.post("/compare-schema")
def compare_schema_route():
    d, source, target = db_paths()
    if not source or not source.exists():
        return redirect(url_for("index"))

    try:
        report = compare_schema(source, target)
        session["schema_report"] = report
    except Exception as exc:
        session["schema_report"] = {
            "matches": False,
            "error": str(exc),
            "tables_added": [],
            "tables_deleted": [],
            "table_mismatches": [],
        }

    return redirect(url_for("dashboard",
        table_filter=request.form.get("table_filter", "all"),
        table_name_filter=request.form.get("table_name_filter", ""),
    ))


@app.route("/dashboard")
def dashboard():
    d, source, target = db_paths()
    if not source or not source.exists():
        return redirect(url_for("index"))

    stats = session.get("stats") or compare_databases(source, target)
    session["stats"] = stats

    # A browser may still have an older session cookie that predates one or
    # more patch options. Merge stored values over defaults so dashboard
    # rendering remains compatible with those sessions.
    options = {
        "change_structure": False,
        "delete_rows": False,
        "update_rows": False,
    }
    stored_options = session.get("patch_options")
    if isinstance(stored_options, dict):
        options.update(stored_options)

    summary = stats["summary"]
    impact = {
        "tables_added": summary["tables_added"] if options["change_structure"] else 0,
        "tables_changed": 0,
        "tables_deleted": 0,
        "columns_added": summary["columns_added"] if options["change_structure"] else 0,
        "columns_changed": 0,
        "columns_deleted": 0,
        "rows_added": summary["rows_added"],
        "rows_changed": summary["rows_changed"] if options["update_rows"] else 0,
        "rows_deleted": summary["rows_deleted"] if options["delete_rows"] else 0,
    }

    # Server-side table filtering and pagination.
    table_filter = request.args.get("table_filter", "all")
    table_name_filter = request.args.get("table_name_filter", "").strip()
    try:
        page = max(1, int(request.args.get("page", "1")))
    except ValueError:
        page = 1

    allowed_filters = {
        "all", "tables_added", "columns_added",
        "rows_added", "rows_changed", "rows_deleted"
    }
    if table_filter not in allowed_filters:
        table_filter = "all"

    filtered_tables = []
    pattern = table_name_filter.lower()
    for table in stats.get("tables", []):
        change_types = set(table.get("change_types", []))
        type_match = table_filter == "all" or table_filter in change_types
        name_match = not pattern or fnmatch.fnmatchcase(
            table.get("name", "").lower(), pattern
        )
        if type_match and name_match:
            filtered_tables.append(table)

    total_tables = len(filtered_tables)
    total_pages = max(1, math.ceil(total_tables / PAGINATION_SIZE))
    page = min(page, total_pages)
    start_index = (page - 1) * PAGINATION_SIZE
    page_tables = filtered_tables[start_index:start_index + PAGINATION_SIZE]

    pagination = {
        "page": page,
        "size": PAGINATION_SIZE,
        "total": total_tables,
        "pages": total_pages,
        "start": start_index + 1 if total_tables else 0,
        "end": min(start_index + PAGINATION_SIZE, total_tables),
        "has_prev": page > 1,
        "has_next": page < total_pages,
    }

    return render_template(
        "dashboard.html",
        stats=stats,
        impact=impact,
        options=options,
        patch_name=session.get("patch_name"),
        base_name=session.get("base_name"),
        merged=session.get("merged", False),
        schema_report=session.get("schema_report"),
        table_rows=page_tables,
        pagination=pagination,
    )


@app.route("/changes")
def changes():
    d, source, target = db_paths()
    if not source or not source.exists():
        return redirect(url_for("index"))

    table_names = sorted(set(
        [x["name"] for x in session.get("stats", {}).get("tables", [])]
    ))
    selected = request.args.get("table")
    detail = None

    if selected and selected in table_names:
        detail = get_table_changes(source, target, selected)

    return render_template(
        "changes.html",
        tables=table_names,
        selected=selected,
        detail=detail,
        table_filter=request.args.get("table_filter", "all"),
        table_name_filter=request.args.get("table_name_filter", ""),
        page=request.args.get("page", "1"),
    )


@app.route("/preview")
def preview():
    d, source, target = db_paths()
    if not source or not source.exists() or not target.exists():
        return redirect(url_for("index"))

    database = request.args.get("database", "patch")
    merged_path = d / "merged.sqlite"
    if database not in {"patch", "base", "merged"}:
        database = "patch"
    if database == "merged" and not merged_path.exists():
        flash("Run Merge Now before previewing the Merged database.", "warning")
        return redirect(url_for("preview", database="base"))
    database_path = {
        "patch": source,
        "base": target,
        "merged": merged_path,
    }[database]
    try:
        page = max(1, int(request.args.get("page", "1")))
    except ValueError:
        page = 1
    connection = sqlite3.connect(database_path)
    try:
        table_names = tables(connection)
    finally:
        connection.close()
    selected = request.args.get("table")
    if selected not in table_names:
        selected = table_names[0] if table_names else None

    page_size = 100
    total = 0
    detail = None
    if selected:
        detail = get_table_preview(database_path, selected, page=page, page_size=page_size)
        total = detail["total"]
        page_count = max(1, math.ceil(total / page_size))
        page = min(page, page_count)
        if detail["page"] != page:
            detail = get_table_preview(database_path, selected, page=page, page_size=page_size)
    else:
        page_count = 1

    return render_template(
        "preview.html",
        database=database,
        database_name={
            "patch": session.get("patch_name"),
            "base": session.get("base_name"),
            "merged": "merged.sqlite",
        }[database],
        tables=table_names,
        selected=selected,
        detail=detail,
        page=page,
        page_count=page_count,
        page_size=page_size,
        total=total,
    )


@app.post("/merge")
def merge():
    d, source, target = db_paths()
    if not source or not source.exists():
        return redirect(url_for("index"))

    change_structure = request.form.get("change_structure") == "yes"
    delete_rows = request.form.get("delete_rows") == "yes"
    update_rows = request.form.get("update_rows") == "yes"

    session["patch_options"] = {
        "change_structure": change_structure,
        "delete_rows": delete_rows,
        "update_rows": update_rows,
    }

    merged = d / "merged.sqlite"
    try:
        merge_databases(source, target, merged, change_structure=change_structure, delete_rows=delete_rows, update_rows=update_rows)
        # Keep the dashboard's detected differences anchored to the original
        # Patch/Base pair. Comparing against merged.sqlite would consume the
        # differences just applied and make the Added counts drop to zero.
        stats = compare_databases(source, target)
        session["stats"] = stats
        session["merged"] = True
        flash("Merge completed. The original Patch and Base databases were not modified.", "success")
    except Exception as exc:
        flash(f"Merge failed: {exc}", "danger")
    return redirect(url_for("dashboard"))


@app.post("/refresh-dashboard")
def refresh_dashboard():
    d, source, target = db_paths()
    if not source or not source.exists():
        return redirect(url_for("index"))

    session["patch_options"] = {
        "change_structure": request.form.get("change_structure", "yes") == "yes",
        "delete_rows": request.form.get("delete_rows") == "yes",
        "update_rows": request.form.get("update_rows") == "yes",
    }

    # Re-read the Patch/Base comparison and redisplay the dashboard.
    session["stats"] = compare_databases(source, target)
    schema_result = compare_schema(source, target)
    session["schema_match"] = schema_result.get("matches", False)
    session["schema_report"] = None
    return redirect(url_for(
        "dashboard",
        table_filter=request.form.get("table_filter", "all"),
        table_name_filter=request.form.get("table_name_filter", ""),
    ))


@app.route("/download")
def download():
    d, source, target = db_paths()
    merged = d / "merged.sqlite" if d else None
    if not merged or not merged.exists():
        flash("No patched database is available. Run Merge Now first.", "warning")
        return redirect(url_for("dashboard"))

    return send_file(
        merged,
        as_attachment=True,
        download_name="merged.sqlite",
        mimetype="application/x-sqlite3",
    )


@app.post("/reset")
def reset():
    d, _, _ = db_paths()
    if d:
        shutil.rmtree(d, ignore_errors=True)
    session.clear()
    return redirect(url_for("index"))


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
